"""
Собирает единый Word-документ из markdown-файлов пакета документации ГОСТ 34
(docs/gost/01-tz.md ... 04-rukovodstvo-administratora.md) для сдачи преподавателю.

Разбирает только тот подмножество markdown, которое реально используется в этих
файлах (заголовки #/##/###, таблицы, **жирный текст**, списки - и 1., блоки ```,
горизонтальные линии ---, ссылки [текст](урл)) - не претендует на полный CommonMark.

Запуск: python build_docx.py
Результат: ../../expenses-gost-documentation.docx (корень репозитория).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GOST_DIR = Path(__file__).parent
ROOT = GOST_DIR.parent.parent
OUT = ROOT / "expenses-gost-documentation.docx"

FONT = "Times New Roman"
BODY_SIZE = Pt(14)

DOCS = [
    ("01-tz.md", None),
    ("02-poyasnitelnaya-zapiska.md", None),
    ("03-rukovodstvo-polzovatelya.md", None),
    ("04-rukovodstvo-administratora.md", None),
]

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def strip_inline(text: str) -> str:
    text = LINK_RE.sub(r"\1", text)
    return text


def add_run_with_bold(paragraph, text: str, size=BODY_SIZE, bold_default=False):
    """Добавляет текст в параграф, разбирая **жирный** внутри строки."""
    text = strip_inline(text)
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.font.name = FONT
            r.font.size = size
            r.bold = bold_default
        r = paragraph.add_run(m.group(1))
        r.font.name = FONT
        r.font.size = size
        r.bold = True
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.font.name = FONT
        r.font.size = size
        r.bold = bold_default


def set_cell_text(cell, text, bold=False, size=Pt(12)):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(strip_inline(text).strip())
    r.font.name = FONT
    r.font.size = size
    r.bold = bold


def add_page_break(doc):
    doc.add_page_break()


def add_heading(doc, text, level):
    text = strip_inline(text).strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    size = {1: Pt(18), 2: Pt(16), 3: Pt(14)}.get(level, Pt(14))
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = size
    r.bold = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_table(doc, rows):
    """rows: список списков ячеек (первая строка — заголовок)."""
    n_cols = len(rows[0])
    table = doc.add_table(rows=0, cols=n_cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            if j < n_cols:
                set_cell_text(cells[j], val, bold=(i == 0))
    doc.add_paragraph()


def render_markdown(doc, md_text: str):
    lines = md_text.split("\n")
    i = 0
    n = len(lines)
    in_code = False
    code_buf = []
    while i < n:
        line = lines[i]

        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(11)
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        # Заголовки
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            add_heading(doc, m.group(2), level)
            i += 1
            continue

        # Таблицы: строка с | ... |, следующая — разделитель |---|---|
        if line.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            table_lines = [line]
            j = i + 2
            while j < n and lines[j].strip().startswith("|"):
                table_lines.append(lines[j])
                j += 1
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                rows.append(cells)
            add_table(doc, rows)
            i = j
            continue

        # Списки (- ... или N. ...)
        m_bul = re.match(r"^\s*-\s+(.*)$", line)
        m_num = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m_bul or m_num:
            p = doc.add_paragraph(style="List Bullet" if m_bul else "List Number")
            add_run_with_bold(p, (m_bul or m_num).group(1))
            i += 1
            continue

        # Обычный абзац — собираем до пустой строки
        para_lines = [line]
        j = i + 1
        while j < n and lines[j].strip() and not lines[j].strip().startswith(("#", "|", "-", "```")) \
                and not re.match(r"^\s*\d+\.\s+", lines[j]):
            para_lines.append(lines[j])
            j += 1
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        add_run_with_bold(p, " ".join(s.strip() for s in para_lines))
        i = j


def set_default_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = BODY_SIZE
    style.paragraph_format.line_spacing = 1.5
    # Кириллица иногда требует явного указания шрифта в rFonts/eastAsia
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT)

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)


def add_title_page(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("[НАИМЕНОВАНИЕ УЧЕБНОГО ЗАВЕДЕНИЯ]")
    r.font.name = FONT
    r.font.size = Pt(14)

    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("КУРСОВАЯ РАБОТА")
    r.font.name = FONT
    r.font.size = Pt(20)
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("по дисциплине «[НАЗВАНИЕ ДИСЦИПЛИНЫ]»")
    r.font.name = FONT
    r.font.size = Pt(14)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Разработка клиент-серверной модели приложения учёта личных расходов «Expenses»")
    r.font.name = FONT
    r.font.size = Pt(16)
    r.bold = True

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for line in ["Выполнил: [ФИО студента], группа [___]", "Проверил: [ФИО преподавателя]"]:
        r = p.add_run(line)
        r.font.name = FONT
        r.font.size = Pt(14)
        p.add_run().add_break()

    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026")
    r.font.name = FONT
    r.font.size = Pt(14)

    doc.add_page_break()


def add_toc(doc, titles):
    add_heading(doc, "Содержание", 1)
    for title in titles:
        p = doc.add_paragraph()
        add_run_with_bold(p, title)
    doc.add_page_break()


def main():
    doc = Document()
    set_default_style(doc)
    add_title_page(doc)

    titles = [
        "1. Техническое задание",
        "2. Пояснительная записка к техническому проекту",
        "3. Руководство пользователя",
        "4. Руководство администратора",
    ]
    add_toc(doc, titles)

    for idx, (filename, _) in enumerate(DOCS):
        path = GOST_DIR / filename
        md_text = path.read_text(encoding="utf-8")
        render_markdown(doc, md_text)
        if idx < len(DOCS) - 1:
            add_page_break(doc)

    doc.save(OUT)
    print(f"Готово: {OUT}")


if __name__ == "__main__":
    main()
